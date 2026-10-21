"""Build academic HTML/PPTX assets offline. No runtime, API or database access."""
from pathlib import Path
import json, html, re
import markdown
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
FONT='DejaVu Sans'
COL={'navy':'112B3C','teal':'007D7A','paper':'F5F6F2','white':'FFFFFF','ink':'183B4A','muted':'536E79','line':'DAE4E3','amber':'AF5F22'}
def esc(s):return html.escape(str(s))
def status_text():
 s=json.loads((BASE/'delivery-status.json').read_text())
 return s, '\n'.join(f'- **{key}:** {value if value is not None else "Chưa có xác nhận"}' for key,value in s.items())
s,status=status_text()
text=(BASE/'Final_Project_Report.md').read_text()
# Add the changing release evidence as a separate cohort; never edit Phase 4 artifacts.
appendix='\n\n## Phụ lục cập nhật phát hành ngày 21/10/2026\n\nMốc kết quả trong các chương là bằng chứng Phase 4 và các cohort lịch sử được ghi rõ. Trạng thái sau mốc này được tách riêng dưới đây; không ghi đè S06/C01 hoặc kết quả Phase 3A.\n\n'+status+'\n\nCác trạng thái không có receipt chưa được coi là PASS. [Dữ liệu trạng thái](delivery-status.json); [hướng dẫn demo](DEMO_GUIDE.md).\n'
md=markdown.Markdown(extensions=['tables','fenced_code','toc','sane_lists'],extension_configs={'toc':{'toc_depth':'2-3','title':''}})
body=md.convert(text+appendix)
body=body[body.index('<h2'):]
cover='''<section class="cover"><div class="school">TRƯỜNG ĐẠI HỌC PHENIKAA</div><div class="cover-rule"></div><div class="eyebrow">BÁO CÁO ĐỒ ÁN CUỐI KỲ</div><h1>Hanoi Land Law<br>Assistant</h1><p class="subtitle">Hệ thống tra cứu pháp luật về bồi thường, hỗ trợ và tái định cư cho Vành đai 4 tại Hà Nội bằng RAG có kiểm soát nguồn</p><table class="cover-meta"><tr><td>Khoa / viện</td><td>____________________________</td></tr><tr><td>Học phần / mã học phần</td><td>____________________________</td></tr><tr><td>Sinh viên / MSSV</td><td>____________________________</td></tr><tr><td>Lớp / khóa</td><td>____________________________</td></tr><tr><td>Giảng viên hướng dẫn</td><td>____________________________</td></tr></table><p class="cover-date">Hà Nội · 21 tháng 10 năm 2026</p><p class="cover-note">Phạm vi tra cứu có nguồn; không xác định quyền lợi cá nhân.<br>Thông tin hành chính để trống theo yêu cầu; chưa có rubric chính thức.</p></section>'''
css='''@page { size:A4; margin:20mm 18mm 20mm 22mm; } *{box-sizing:border-box} body{font:11pt/1.58 "DejaVu Serif",serif;color:#1b2c35;margin:0} h1,h2,h3,h4{font-family:"DejaVu Sans",sans-serif;color:#112b3c;line-height:1.28} h2{font-size:18pt;margin:25px 0 14px;break-after:avoid} h3{font-size:12pt;margin:18px 0 9px;break-after:avoid} p{margin:8px 0;text-align:justify;orphans:3;widows:3} a{color:#00736f;text-decoration:none;overflow-wrap:anywhere} strong{font-weight:700} table{border-collapse:collapse;width:100%;font:8.4pt/1.43 "DejaVu Sans",sans-serif;margin:12px 0;table-layout:fixed;break-inside:auto} th{background:#e8f1ef;color:#112b3c;text-align:left} td,th{border:1px solid #cbd8d8;padding:7px 8px;vertical-align:top;overflow-wrap:anywhere} tr{break-inside:avoid} thead{display:table-header-group} code{font:8.5pt/1.45 "DejaVu Sans Mono",monospace;overflow-wrap:anywhere;background:#f0f3f3} pre{white-space:pre-wrap;overflow-wrap:anywhere;padding:10px;background:#f0f3f3;break-inside:avoid} li{margin:6px 0} img{max-width:100%;height:auto;break-inside:avoid} .cover{height:248mm;break-after:page;text-align:center;font-family:"DejaVu Sans";padding-top:12mm} .school{font-size:15pt;font-weight:700;letter-spacing:.3px}.cover-rule{height:3px;background:#007d7a;width:70mm;margin:14mm auto 20mm}.eyebrow{font-size:12pt;letter-spacing:2px}.cover h1{font-size:35pt;line-height:1.16;margin:10mm 0 7mm}.subtitle{text-align:center;font-size:13pt;line-height:1.55;margin:0 4mm 9mm}.cover-meta{width:145mm;margin:10mm auto;font-size:10pt}.cover-meta td{border:0;text-align:left;padding:6px}.cover-meta td:first-child{width:58mm}.cover-date{text-align:center;margin-top:17mm;font-size:12pt}.cover-note{text-align:center;font-size:8pt;color:#667c82;margin-top:8mm}.toc{font-family:"DejaVu Sans";font-size:10pt;break-after:page}.toc ul{list-style:none;padding-left:13px}.toc>ul{padding:0}.toc li{margin:5px 0}.toc-title{font-size:20pt;font-weight:700;color:#112b3c;margin-bottom:18px} h2[id^="1-"]{break-before:page} h2[id^="2-"],h2[id^="3-"],h2[id^="4-"],h2[id^="5-"],h2[id^="6-"],h2[id^="7-"],h2[id^="8-"],h2[id^="9-"],h2[id^="10-"],h2[id^="11-"],h2[id^="12-"]{break-before:page}'''
(BASE/'Final_Project_Report.html').write_text('<!doctype html><html lang="vi"><meta charset="utf-8"><title>Hanoi Land Law Assistant — Báo cáo đồ án</title><style>'+css+'</style><body>'+cover+'<section class="toc"><div class="toc-title">Mục lục</div>'+md.toc+'</section><main>'+body+'</main></body></html>')
slides=json.loads((BASE/'slides.json').read_text())
if not (BASE/'architecture.png').exists():
 raise SystemExit('Run node export_pdf.cjs --prepare-assets-only before building PPTX.')
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
def rgb(name):return RGBColor.from_string(COL.get(name,name))
def rect(sl,x,y,w,h,fill,line=None,radius=False):
 shape=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill);shape.line.fill.background()
 if radius:
  shape.adjustments[0]=.12
 return shape
def box(sl,x,y,w,h,txt,size=22,color='ink',bold=False):
 sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w), Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for j,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(6)
 return sh
def cards(sl,items,y=4.15,height=1.65):
 n=len(items);gap=.22;w=(11.73-gap*(n-1))/n
 for j,(label,desc) in enumerate(items):
  x=.8+j*(w+gap);rect(sl,x,y,w,height,'white',radius=True);rect(sl,x,y,.05,height,'teal');box(sl,x+.2,y+.16,w-.4,.35,label,11,'teal',True);box(sl,x+.2,y+.6,w-.4,height-.64,desc,17 if n==3 else 16,'ink')
for i,data in enumerate(slides):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);rect(sl,0,0,13.3333,7.5,'paper');rect(sl,0,0,13.3333,.09,'teal')
 box(sl,.8,.38,11.7,.35,data['kicker'],11,'teal',True)
 kind=data.get('kind','normal')
 if kind=='cover':
  box(sl,.8,1.12,11.6,1.8,data['title'],46,'navy',True);box(sl,.82,3.1,11.5,.75,data['lead'],23,'muted');cards(sl,data['cards'],4.35,1.7)
 elif kind=='architecture':
  box(sl,.8,1,11.7,.68,data['title'],33,'navy',True);box(sl,.8,1.85,11.7,.65,data['lead'],18,'muted');img=BASE/'architecture.png'
  if img.exists(): sl.shapes.add_picture(str(img),Inches(2.95),Inches(2.65),height=Inches(3.65))
  else: box(sl,.8,3,11.7,2,'Xem docs/diagrams/architecture.svg trong source.',22)
 elif kind=='metrics':
  box(sl,.8,1,11.7,1.15,data['title'],31,'navy',True);box(sl,.8,2.35,11.7,.55,data['lead'],18,'muted')
  for j,(val,label) in enumerate(data['metrics']):
   x=.8+j*4;box(sl,x,3.12,3.8,.7,val,39,'teal',True);box(sl,x,3.94,3.8,.4,label,13,'muted')
  cards(sl,data['cards'],4.75,1.3)
 else:
  box(sl,.8,1.05,11.7,1.45,data['title'],34,'navy',True);box(sl,.8,2.78,11.7,.95,data['lead'],20,'muted');cards(sl,data.get('cards',[]),4.25,1.75)
 box(sl,.8,6.75,11.2,.3,data['evidence'],9,'muted');box(sl,12.1,6.7,.5,.35,f'{i+1:02}',11,'teal',True)
 sl.notes_slide.notes_text_frame.text=data['notes']+'\n\nBằng chứng: '+data['evidence']+'\nTrạng thái phát hành/pilot cập nhật: '+json.dumps(s,ensure_ascii=False)
prs.core_properties.title='Hanoi Land Law Assistant — Đồ án liên ngành';prs.core_properties.subject='Tra cứu pháp luật có nguồn cho Vành đai 4 Hà Nội';prs.core_properties.author='';prs.core_properties.keywords='RAG, legal source, Vietnamese, academic prototype'
prs.save(BASE/'Final_Project_Presentation.pptx')
slidecss='''@page{size:13.333333in 7.5in;margin:0}*{box-sizing:border-box}body{margin:0;font-family:"DejaVu Sans",sans-serif;color:#183b4a}.slide{width:1280px;height:720px;padding:36px 77px 40px;background:#f5f6f2;border-top:9px solid #007d7a;position:relative;break-after:page;overflow:hidden}.slide:last-child{break-after:auto}.kicker{font-size:14px;color:#007d7a;font-weight:700;letter-spacing:1px}h1{white-space:pre-line;font-size:44px;color:#112b3c;line-height:1.2;margin:48px 0 24px;max-width:1110px} .lead{font-size:25px;line-height:1.45;color:#536e79;margin:0;min-height:90px}.cards{display:flex;gap:21px;margin-top:32px}.card{flex:1;background:white;border-left:5px solid #007d7a;border-radius:9px;padding:18px 20px;height:168px}.card-label{font-size:14px;font-weight:bold;color:#007d7a;margin-bottom:18px}.card-desc{font-size:21px;line-height:1.42;white-space:pre-line}.footer{position:absolute;left:77px;bottom:44px;font-size:11px;color:#536e79;max-width:1040px}.page{position:absolute;right:77px;bottom:41px;font-size:15px;color:#007d7a}.cover h1{font-size:62px;margin-top:45px}.cover .lead{font-size:29px;min-height:62px}.cover .cards{margin-top:31px}.cover .card{height:165px}.cover .card-desc{font-size:20px}.architecture h1{font-size:42px;margin:25px 0 12px}.architecture .lead{font-size:22px;min-height:0}.architecture img{display:block;width:560px;height:350px;object-fit:contain;margin:16px auto 0}.metrics h1{font-size:40px;margin:28px 0 15px}.metrics .lead{font-size:22px;min-height:0}.metric-row{display:flex;gap:25px;margin-top:27px}.metric{width:33.33%}.metric-value{font-size:52px;color:#007d7a;font-weight:bold}.metric-label{font-size:17px;color:#536e79;margin-top:5px}.metrics .cards{margin-top:30px}.metrics .card{height:125px}.metrics .card-label{margin-bottom:12px}.metrics .card-desc{font-size:19px}'''
parts=[]
for i,d in enumerate(slides):
 kind=d.get('kind','normal');cardshtml=''.join('<div class="card"><div class="card-label">'+esc(label)+'</div><div class="card-desc">'+esc(desc)+'</div></div>' for label,desc in d.get('cards',[]))
 extra='<img src="../../docs/diagrams/architecture.svg" alt="Sơ đồ kiến trúc được tái sử dụng từ repository">' if kind=='architecture' else ''
 if kind=='metrics': extra='<div class="metric-row">'+''.join('<div class="metric"><div class="metric-value">'+esc(v)+'</div><div class="metric-label">'+esc(l)+'</div></div>' for v,l in d['metrics'])+'</div>'
 parts.append(f'<section class="slide {kind}" data-slide="{i+1}"><div class="kicker">{esc(d["kicker"])}</div><h1>{esc(d["title"])}</h1><p class="lead">{esc(d["lead"])}</p>{extra}<div class="cards">{cardshtml}</div><div class="footer">{esc(d["evidence"])}</div><div class="page">{i+1:02}</div></section>')
(BASE/'Final_Project_Presentation.html').write_text('<!doctype html><html lang="vi"><meta charset="utf-8"><title>Hanoi Land Law Assistant — Presentation</title><style>'+slidecss+'</style><body>'+''.join(parts)+'</body></html>')
(BASE/'SPEAKER_NOTES.md').write_text('# Ghi chú thuyết trình — 12 slide\n\n'+ '\n\n'.join(f'## {i+1}. {d["title"].replace(chr(10), " ")}\n\n{d["notes"]}\n\nEvidence: {d["evidence"]}' for i,d in enumerate(slides)))
print('Built report HTML, presentation HTML/PPTX, and speaker notes.')
